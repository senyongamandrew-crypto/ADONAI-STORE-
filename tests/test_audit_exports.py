"""Downloadable CSV / genuine Excel regressions — throwaway SQL only.
Run separately: .venv/bin/python -m unittest discover -s tests -p test_audit_exports.py -v
"""
import csv
from datetime import datetime, timedelta
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook
from sqlalchemy import event

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_TEMP = tempfile.TemporaryDirectory(prefix="adonai-export-tests-")
os.environ.update({"DATABASE_URL":"sqlite:///" + str(Path(_TEMP.name) / "exports.db"),
                   "ADONAI_REQUIRE_STAFF_AUTH":"0", "ADONAI_RATE_LIMIT_DISABLED":"1",
                   "ADONAI_SEED_DEMO_CATALOG":"0", "ADONAI_SEED_DEMO_USERS":"0"})
import database
database.IS_POSTGRES = False
from database import Base, engine, get_db
from models import (ExportDownload, Order, OrderItem, Product, User, LedgerEntry,
                    ExpenseRecord, AccountingJournalEntry, AccountingJournalLine,
                    PosAccessKey, PosAccessSession, StoreSetting)
import access_control
import api
import exports
import serve

AT_START = datetime(2026, 10, 9, 21, 0, 0)  # October 10, 00:00 EAT
AT_END = datetime(2026, 10, 10, 20, 59, 59)
NAME = 'Nakyeyune, "Sarah"\nKampala'


class AuditExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): Base.metadata.create_all(engine)

    def setUp(self):
        with get_db() as session:
            for model in [ExportDownload, OrderItem, Order, LedgerEntry, ExpenseRecord,
                          AccountingJournalLine, AccountingJournalEntry, Product, User,
                          PosAccessSession, PosAccessKey, StoreSetting]:
                session.query(model).delete(synchronize_session=False)
            session.add(Product(id="P1", sku="SKU-001", barcode_id="BC-001", name="Floral dress", demographic="Women", category="Dresses", cost_price=10000, selling_price=20000, base_price=20000, in_stock_count=3, created_at=AT_START))
            session.add(User(id="STF-EXPORT", name="Fixture cashier", phone="+256700111222", role="cashier", pin_hash="NEVER-EXPORT-THIS-PIN-HASH", active=True, created_at=AT_START))
            self.order(session, "OLD", AT_START - timedelta(seconds=1))
            self.order(session, "START", AT_START, created=AT_START - timedelta(days=2))
            self.order(session, "END", AT_END)
            self.order(session, "NEXT", AT_END + timedelta(seconds=1))
            self.order(session, "UNPAID", AT_START, payment="unpaid", channel="web")
            self.order(session, "CANCELLED", AT_START, status="cancelled")
            for identity, kind, amount in [("L-SALE", "sale", 43000), ("L-EXP", "expense", -5000), ("L-REF", "refund", -2000), ("L-ADJ", "adjustment", 1000)]:
                session.add(LedgerEntry(id=identity, kind=kind, amount=amount, category="Cash", description=kind + " posting", ref_id="START", staff_name="Fixture cashier", created_at=AT_START))
            session.add(ExpenseRecord(id="EXP1", category="Rent", amount=5000, payment_method="cash", vendor="Landlord", receipt_reference="R-1", status="POSTED", occurred_at=AT_START, created_by_name="Accounts"))
            entry = AccountingJournalEntry(id="J1", reference="EXP1", source="EXPENSE", description="Rent", status="POSTED", occurred_at=AT_START, created_by_name="Accounts")
            entry.lines = [AccountingJournalLine(id="JL1", account_code="5100", account_name="Rent", debit=5000, credit=0), AccountingJournalLine(id="JL2", account_code="1000", account_name="Cash", debit=0, credit=5000)]
            session.add(entry)

    def order(self, session, identity, paid_date, payment="paid", status="completed", created=None, channel="pos"):
        order = Order(id=identity, channel=channel, status=status, dispatch_status="Completed", customer_name=NAME, customer_phone="+256700123456", customer_address="Makerere, Kampala", delivery_area="Central", delivery_type="boda", delivery_notes="Call at gate", subtotal=40000, delivery_fee=3000, total=43000, tender_type="cash", payment_status=payment, paid_at=paid_date if payment == "paid" else None, cashier_name="Fixture cashier", created_at=created or paid_date)
        order.items = [OrderItem(id="IT-" + identity, name="Floral dress", qty=2, unit_price=20000, unit_cost=10000, line_total=40000, product_id="P1")]
        session.add(order)

    def raw(self, method, path, body=None, headers=None):
        return api.handle_api_request(method, path, {}, json.dumps(body).encode() if body is not None else b"", headers or {}, "export-test")

    def prepare(self, kind="sales", file_format="csv", headers=None, **extra):
        result = self.raw("POST", "/api/exports/prepare", {"kind":kind, "format":file_format, **extra}, headers)
        self.assertEqual(result[0], 201, result[2])
        return json.loads(result[2])["download"]

    def csv(self, ticket):
        result = self.raw("GET", ticket["url"])
        self.assertEqual(result[0], 200, result[2]); self.assertTrue(result[2].startswith(b"\xef\xbb\xbf"))
        return list(csv.DictReader(io.StringIO(result[2].decode("utf-8-sig"), newline="")))

    def test_csv_is_a_real_file_with_headers_and_inclusive_eat_paid_dates(self):
        ticket = self.prepare(**{"from":"2026-10-10", "to":"2026-10-10"})
        self.assertTrue(ticket["filename"].endswith(".csv")); self.assertEqual(ticket["records"], 2)
        rows = self.csv(ticket)
        self.assertEqual({r["Receipt / Order ID"] for r in rows}, {"START", "END"})
        self.assertEqual(rows[0]["Customer"], NAME)
        self.assertEqual(rows[0]["Phone"], "'+256700123456")
        self.assertEqual(rows[0]["Total (UGX)"], "43000")
        self.assertEqual(rows[0]["Posted / Order date (EAT)"], "2026-10-10 00:00:00 +03:00")
        self.assertIn("@ 20000 UGX", rows[0]["Items"])
        result = self.raw("GET", ticket["url"])
        self.assertEqual(len(result[2]), ticket["bytes"])
        self.assertIn("attachment;", result.headers["Content-Disposition"])
        self.assertIn(ticket["filename"], result.headers["Content-Disposition"])
        self.assertEqual(result.headers["Cache-Control"], "private, no-store")
        self.assertEqual(result.headers["Referrer-Policy"], "no-referrer")
        self.assertEqual(ticket["timezone"], "Africa/Kampala")

    def test_completed_pos_legacy_payment_flag_is_not_lost_and_unpaid_web_is_excluded(self):
        with get_db() as session:
            self.order(session, "LEGACY-POS", AT_START, payment="unpaid", channel="pos")
        ticket = self.prepare(**{"from":"2026-10-10", "to":"2026-10-10"})
        self.assertEqual({row["Receipt / Order ID"] for row in self.csv(ticket)}, {"START", "END", "LEGACY-POS"})

    def test_financial_ledger_is_actual_signed_postings_not_sales_copies(self):
        ticket = self.prepare("ledger", **{"from":"2026-10-10", "to":"2026-10-10"})
        rows = self.csv(ticket)
        self.assertEqual({row["Ledger ID"] for row in rows}, {"L-SALE", "L-EXP", "L-REF", "L-ADJ"})
        self.assertEqual(sum(int(row["Amount (UGX, signed)"]) for row in rows), 37000)
        self.assertIn("-5000", [row["Amount (UGX, signed)"] for row in rows])
        filtered = self.prepare("ledger", q="refund"); self.assertEqual(filtered["records"], 1)
        self.assertEqual(self.csv(filtered)[0]["Type"], "refund")

    def test_xlsx_is_a_genuine_workbook_numeric_amounts_and_literal_text(self):
        with get_db() as session: session.get(Order, "START").customer_name = " \t=1+1"
        ticket = self.prepare("sales", "xlsx", q="START")
        result = self.raw("GET", ticket["url"])
        self.assertEqual(result[0], 200); self.assertTrue(result[2].startswith(b"PK"))
        self.assertEqual(result[1], exports.MIMES["xlsx"])
        book = load_workbook(io.BytesIO(result[2]), data_only=False)
        sheet = book.active; cells = list(sheet.rows)
        self.assertEqual(len(cells), 2)
        self.assertEqual(cells[1][2].value, " \t=1+1"); self.assertEqual(cells[1][2].data_type, "s")
        self.assertEqual(cells[1][3].value, "+256700123456"); self.assertEqual(cells[1][3].number_format, "@")
        self.assertEqual(cells[1][-1].value, 43000); self.assertEqual(cells[1][-1].data_type, "n")
        self.assertEqual(sheet.freeze_panes, "A2"); self.assertTrue(sheet.auto_filter.ref)
        book.close()
        ledger = self.prepare("ledger", "xlsx", q="L-EXP")
        book = load_workbook(io.BytesIO(self.raw("GET", ledger["url"])[2]))
        self.assertEqual(book.active.cell(2, 6).value, -5000); book.close()

    def test_orders_include_all_statuses_and_contacts_and_customers_are_real(self):
        orders = self.prepare("orders", **{"from":"2026-10-10", "to":"2026-10-10"})
        rows = self.csv(orders)
        self.assertEqual({row["Receipt / Order ID"] for row in rows}, {"END", "UNPAID", "CANCELLED"})
        self.assertEqual(rows[0]["Delivery notes"], "Call at gate")
        customers = self.prepare("customers", **{"from":"2026-10-10", "to":"2026-10-10"})
        contacts = self.csv(customers)
        self.assertEqual(len(contacts), 1); self.assertEqual(contacts[0]["Customer"], NAME)
        self.assertEqual(contacts[0]["Orders in period"], "3"); self.assertEqual(contacts[0]["Paid spend (UGX)"], "43000")
        self.assertEqual(self.prepare("customers", q="Makerere")["records"], 1)

    def test_other_reports_empty_dates_and_no_secrets_or_mutations(self):
        with get_db() as session:
            before = (session.query(Order).count(), session.query(LedgerEntry).count(), session.get(Product, "P1").in_stock_count)
        for kind, count in [("expenses", 1), ("journal", 2), ("stock", 1), ("staff", 1)]:
            for file_format in ("csv", "xlsx"):
                ticket = self.prepare(kind, file_format, **{"from":"2026-10-10", "to":"2026-10-10"})
                self.assertEqual(ticket["records"], count)
                result = self.raw("GET", ticket["url"]); self.assertEqual(result[0], 200)
                if file_format == "csv": self.assertNotIn("NEVER-EXPORT", result[2].decode("utf-8-sig"))
        empty = self.prepare(**{"from":"2000-01-01", "to":"2000-01-02"})
        self.assertEqual(empty["records"], 0); self.assertEqual(self.csv(empty), [])
        with get_db() as session:
            self.assertEqual(before, (session.query(Order).count(), session.query(LedgerEntry).count(), session.get(Product, "P1").in_stock_count))

    def test_invalid_options_are_rejected_without_fake_files(self):
        cases = [{"kind":"demo"}, {"kind":{}}, {"format":"xls"}, {"format":"pdf"}, {"from":"2026-02-30"}, {"from":"2026-10-11", "to":"2026-10-10"}, {"to":"9999-12-31"}, {"from":22}]
        for body in cases:
            result = self.raw("POST", "/api/exports/prepare", body)
            self.assertEqual(result[0], 400, (body, result))
        with get_db() as session: self.assertEqual(session.query(ExportDownload).count(), 0)

    def test_head_never_loads_binary_and_http_adapters_preserve_attachment(self):
        ticket = self.prepare(); statements = []
        def collect(connection, cursor, statement, parameters, context, executemany): statements.append(statement)
        event.listen(engine, "before_cursor_execute", collect)
        try: result = self.raw("HEAD", ticket["url"])
        finally: event.remove(engine, "before_cursor_execute", collect)
        self.assertEqual(result[0], 200); self.assertEqual(result[2], b"")
        self.assertEqual(result.headers["Content-Length"], str(ticket["bytes"]))
        self.assertFalse(any("export_downloads.data" in statement.lower() for statement in statements))
        received = {}
        def start(status, headers): received.update(status=status, headers=headers)
        body = b"".join(serve.wsgi_app({"PATH_INFO":ticket["url"], "REQUEST_METHOD":"HEAD", "wsgi.input":io.BytesIO(), "QUERY_STRING":""}, start))
        self.assertEqual(body, b""); self.assertEqual(received["status"], "200 OK")
        self.assertEqual(dict(received["headers"])["Content-Length"], str(ticket["bytes"]))
        self.assertEqual([value for name, value in received["headers"] if name.lower() == "referrer-policy"], ["no-referrer"])

    def test_short_lived_capabilities_are_hashed_expire_and_bound_storage(self):
        ticket = self.prepare(); token = ticket["url"].rsplit("/", 1)[-1]
        self.assertEqual(len(token), 43); self.assertNotIn("?", ticket["url"])
        with get_db() as session:
            item = session.get(ExportDownload, hashlib.sha256(token.encode()).hexdigest())
            self.assertIsNotNone(item); self.assertNotEqual(item.token_hash, token)
            item.expires_at = datetime.utcnow() - timedelta(seconds=1)
        self.assertEqual(self.raw("GET", ticket["url"])[0], 410)
        self.assertEqual(self.raw("GET", "/api/exports/download/bad-token")[0], 404)
        for _ in range(23): self.prepare("staff")
        with get_db() as session: self.assertEqual(session.query(ExportDownload).count(), 20)

    def test_optional_scopes_protect_create_and_revalidate_old_links_on_relock(self):
        owner_password = "Fixture-owner-123!"
        result = self.raw("POST", "/api/pos-access/admin", {"password":owner_password})
        self.assertEqual(result[0], 200, result[2])
        owner_token = json.loads(result[2])["session"]["token"]
        owner_headers = {"x-pos-access":owner_token}
        result = self.raw("PATCH", "/api/pos-access/policy", {"protected_areas":["orders", "finance", "stock"]}, owner_headers)
        self.assertEqual(result[0], 200, result[2])
        for kind in ("sales", "orders", "customers", "ledger", "expenses", "journal", "stock", "staff"):
            self.assertEqual(self.raw("POST", "/api/exports/prepare", {"kind":kind})[0], 403)
        ticket = self.prepare("ledger", headers=owner_headers)
        self.assertEqual(self.raw("GET", ticket["url"])[0], 200, "browser receives an attachment without authentication in its URL")
        self.raw("POST", "/api/pos-access/lock", {}, owner_headers)
        self.assertEqual(self.raw("GET", ticket["url"])[0], 403)

    def test_unlocked_finance_never_grants_order_contacts_or_password_management(self):
        with get_db() as session:
            session.add(PosAccessKey(id=access_control.OWNER_ID, name="Owner", password_hash="fixture", areas="[]", is_admin=True))
            session.add(PosAccessKey(id="F", name="Finance", password_hash="fixture", areas='["finance"]', is_admin=False))
            access_control._set_protected(session, {"orders", "finance", "stock"})
            session.flush()
            session.add(PosAccessSession(token_hash=access_control._digest("F" * 43), key_id="F", key_version=1, expires_at=datetime.utcnow() + timedelta(minutes=30)))
        headers = {"x-pos-access":"F" * 43}
        self.assertEqual(self.prepare("ledger", headers=headers)["records"], 4)
        for kind in ("sales", "orders", "customers", "stock", "staff"):
            self.assertEqual(self.raw("POST", "/api/exports/prepare", {"kind":kind}, headers)[0], 403)

    def test_staff_auth_unchanged_and_later_protection_does_not_reuse_open_ticket(self):
        open_ticket = self.prepare()
        with mock.patch.dict(os.environ, {"ADONAI_REQUIRE_STAFF_AUTH":"1", "STAFF_TERMINAL_KEY":"Fixture-staff-key-123", "STORE_MASTER_KEY":"Fixture-admin-key-123"}):
            self.assertEqual(self.raw("POST", "/api/exports/prepare", {})[0], 401)
            self.assertEqual(self.raw("GET", open_ticket["url"])[0], 403)
            staff = {"x-terminal-key":"Fixture-staff-key-123"}
            self.assertEqual(self.prepare("sales", headers=staff)["records"], 4)
            self.assertEqual(self.raw("POST", "/api/exports/prepare", {"kind":"staff"}, staff)[0], 403)
            self.assertEqual(self.prepare("staff", headers={"x-terminal-key":"Fixture-admin-key-123"})["records"], 1)

    def test_formula_safety_and_size_errors_never_silently_truncate(self):
        for value in ("=1+1", "+256700123456", "-1+2", "@SUM(1)", " \t=HYPERLINK()", "\v=1", "\ufeff=1"):
            self.assertTrue(exports.safe_text(value).startswith("'"))
        with mock.patch.object(exports, "MAX_ROWS", 1):
            result = self.raw("POST", "/api/exports/prepare", {"kind":"ledger"})
            self.assertEqual(result[0], 400); self.assertIn(b"smaller date range", result[2])
        with mock.patch.object(exports, "MAX_ROWS", 1):
            result = self.raw("POST", "/api/exports/prepare", {"kind":"ledger", "format":"xlsx"})
            self.assertEqual(result[0], 400)
        with self.assertRaises(ValueError): exports.render_file(["Text"], [["x" * 32768]], "xlsx", "Test")
        with mock.patch.object(exports, "MAX_BYTES", 12):
            result = self.raw("POST", "/api/exports/prepare", {"kind":"sales"})
            self.assertEqual(result[0], 400)
        with get_db() as session: self.assertEqual(session.query(ExportDownload).count(), 0)


if __name__ == "__main__": unittest.main()
