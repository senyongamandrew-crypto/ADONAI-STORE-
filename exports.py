"""Real CSV / XLSX reports from authoritative records, with short-lived downloads.

Only a cryptographically random, expiring file capability appears in a URL —
never staff passwords, JWTs or area tokens. Permissions are checked both at
creation and download, including area re-locks. No fake demo rows / cached totals.
"""
import csv
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import io
import json
import re
import secrets
import urllib.parse

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, and_, or_
from sqlalchemy.orm import selectinload

import access_control
from database import get_db
from media import MediaResponse
from models import (ExportDownload, Order, Product, User, LedgerEntry, ExpenseRecord,
                    AccountingJournalEntry, AccountingJournalLine)

EAT = timezone(timedelta(hours=3), "EAT")
MAX_ROWS = 100000
MAX_BYTES = 32 * 1024 * 1024
LIFETIME = timedelta(minutes=10)
KINDS = {
    "sales": ("Sales records", "orders"),
    "orders": ("Orders and fulfillment", "orders"),
    "customers": ("Customer details", "orders"),
    "ledger": ("Financial ledger", "finance"),
    "expenses": ("Expenses", "finance"),
    "journal": ("General journal", "finance"),
    "stock": ("Current stock inventory", "stock"),
    "staff": ("Staff directory", "security"),
}
MIMES = {"csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
CANCELED = {"cancelled", "canceled", "void", "expired"}


def response(data, status=200):
    return MediaResponse(status, "application/json; charset=utf-8", json.dumps(data).encode())


def options(body):
    kind, file_format = str(body.get("kind", "sales")), str(body.get("format", "csv"))
    if kind not in KINDS or file_format not in MIMES:
        raise ValueError("Choose a supported report and CSV or Excel (.xlsx).")
    dates = []
    for name in ("from", "to"):
        value = body.get(name) or ""
        if value and (not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value)):
            raise ValueError("Choose valid start and end dates.")
        try: dates.append(date.fromisoformat(value) if value else None)
        except ValueError as exc: raise ValueError("Choose valid start and end dates.") from exc
    start, end = dates
    if start and end and start > end:
        raise ValueError("The start date must be on or before the end date.")
    try:
        lower = datetime.combine(start, time.min, EAT).astimezone(timezone.utc).replace(tzinfo=None) if start else None
        upper = datetime.combine(end + timedelta(days=1), time.min, EAT).astimezone(timezone.utc).replace(tzinfo=None) if end else None
    except (ValueError, OverflowError) as exc: raise ValueError("Choose a supported date range.") from exc
    search = str(body.get("q") or "").strip()[:160].casefold()
    return kind, file_format, lower, upper, start, end, search


def local_date(value):
    if not value: return ""
    if value.tzinfo is None: value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(EAT).strftime("%Y-%m-%d %H:%M:%S +03:00")


def scoped(query, column, lower, upper):
    if lower: query = query.filter(column >= lower)
    if upper: query = query.filter(column < upper)
    return query


def paid(order):
    if (order.status or "").lower() in CANCELED: return False
    status = (order.payment_status or "").lower()
    # The existing POS checkout endpoint (and legacy SQL rows) records
    # completed, paid-at-the-register sales with an old "unpaid" default. Do
    # not discard those real receipts; never treat an unpaid web order as paid.
    return status == "paid" or (order.channel == "pos" and (order.status or "").lower() == "completed" and status in {"", "unpaid"})


def matches(row, search):
    return not search or search in " ".join(str(value or "") for value in row).casefold()


def report_rows(session, kind, lower, upper, search):
    """Return header and an iterator; dates and numbers remain unambiguous."""
    if kind in {"sales", "orders", "customers"}:
        column = func.coalesce(Order.paid_at, Order.created_at) if kind == "sales" else Order.created_at
        query = session.query(Order)
        if kind != "customers": query = query.options(selectinload(Order.items))
        query = scoped(query, column, lower, upper).order_by(column, Order.id)
        if kind == "sales":
            query = query.filter(func.lower(func.coalesce(Order.status, "")).notin_(CANCELED), or_(
                func.lower(func.coalesce(Order.payment_status, "")) == "paid",
                and_(Order.channel == "pos", func.lower(Order.status) == "completed", func.lower(func.coalesce(Order.payment_status, "")).in_(["", "unpaid"]))))
        if kind == "customers":
            header = ["Customer", "Phone", "Address", "Delivery area", "Orders in period", "Paid spend (UGX)", "First order (EAT)", "Latest order (EAT)"]
            def customers():
                groups = {}; scanned = 0
                for order in query.yield_per(500):
                    scanned += 1
                    if scanned > MAX_ROWS: raise ValueError("Too many orders for one export. Choose a smaller date range.")
                    name, phone, address = (order.customer_name or "").strip(), (order.customer_phone or "").strip(), (order.customer_address or "").strip()
                    if not (phone or address or (name and name.casefold() not in {"walk-in", "walk-in customer", "walk in customer"})): continue
                    digits = re.sub(r"\D", "", phone)
                    key = ("phone", digits or phone.casefold()) if phone else ("name/address", name.casefold(), address.casefold())
                    item = groups.setdefault(key, [name, phone, address, order.delivery_area or "", 0, 0, order.created_at, order.created_at])
                    item[0:4] = [name or item[0], phone or item[1], address or item[2], order.delivery_area or item[3]]
                    item[4] += 1; item[5] += int(order.total or 0) if paid(order) else 0
                    item[7] = order.created_at
                for item in groups.values():
                    row = item[:6] + [local_date(item[6]), local_date(item[7])]
                    if matches(row, search): yield row
            return header, customers()
        header = ["Receipt / Order ID", "Posted / Order date (EAT)", "Customer", "Phone", "Address", "Channel", "Payment status", "Payment method", "Order status", "Fulfillment", "Cashier", "Items", "Units", "Subtotal (UGX)", "Delivery (UGX)", "Total (UGX)"]
        if kind == "orders":
            header += ["Delivery type", "Delivery area", "Delivery notes", "Assigned rider", "Payment reference", "Paid date (EAT)", "Updated date (EAT)"]
        def orders():
            for order in query.yield_per(500):
                if kind == "sales" and not paid(order): continue
                row = [order.id, local_date((order.paid_at or order.created_at) if kind == "sales" else order.created_at),
                    order.customer_name or "Walk-in customer", order.customer_phone or "", order.customer_address or "",
                    order.channel, order.payment_status or "", order.tender_type or "", order.status, order.dispatch_status,
                    order.cashier_name or "", "; ".join(f"{item.name} × {item.qty} @ {int(item.unit_price or 0)} UGX" for item in order.items),
                    sum(int(item.qty or 0) for item in order.items), int(order.subtotal or 0), int(order.delivery_fee or 0), int(order.total or 0)]
                if kind == "orders": row += [order.delivery_type or "", order.delivery_area or "", order.delivery_notes or "", order.rider_name or "", order.payment_reference or "", local_date(order.paid_at), local_date(order.updated_at)]
                if matches(row, search): yield row
        return header, orders()
    if kind == "ledger":
        header = ["Ledger ID", "Posting date (EAT)", "Type", "Category", "Description", "Amount (UGX, signed)", "Reference", "Posted by"]
        query = scoped(session.query(LedgerEntry), LedgerEntry.created_at, lower, upper).order_by(LedgerEntry.created_at, LedgerEntry.id)
        iterator = ([row.id, local_date(row.created_at), row.kind, row.category or "", row.description or "", int(row.amount or 0), row.ref_id or "", row.staff_name or ""] for row in query.yield_per(500))
    elif kind == "expenses":
        header = ["Expense ID", "Expense date (EAT)", "Category", "Amount (UGX)", "Paid from", "Vendor / Payee", "Receipt reference", "Status", "Notes", "Posted by"]
        query = scoped(session.query(ExpenseRecord), ExpenseRecord.occurred_at, lower, upper).order_by(ExpenseRecord.occurred_at, ExpenseRecord.id)
        iterator = ([row.id, local_date(row.occurred_at), row.category, int(row.amount), row.payment_method, row.vendor or "", row.receipt_reference or "", row.status, row.notes or "", row.created_by_name or ""] for row in query.yield_per(500))
    elif kind == "journal":
        header = ["Journal ID", "Posting date (EAT)", "Reference", "Source", "Status", "Description", "Account code", "Account", "Debit (UGX)", "Credit (UGX)", "Posted by"]
        query = session.query(AccountingJournalEntry, AccountingJournalLine).join(AccountingJournalLine, AccountingJournalLine.journal_entry_id == AccountingJournalEntry.id)
        query = scoped(query, AccountingJournalEntry.occurred_at, lower, upper).order_by(AccountingJournalEntry.occurred_at, AccountingJournalEntry.id, AccountingJournalLine.id)
        iterator = ([entry.id, local_date(entry.occurred_at), entry.reference, entry.source, entry.status, entry.description or "", line.account_code, line.account_name, int(line.debit or 0), int(line.credit or 0), entry.created_by_name or ""] for entry, line in query.yield_per(500))
    elif kind == "stock":
        header = ["SKU", "Barcode", "Garment", "Category", "Size", "Color", "Condition", "Current quantity", "Current price (UGX)", "Inventory status", "Intake date (EAT)"]
        query = scoped(session.query(Product), Product.created_at, lower, upper).order_by(Product.created_at, Product.id)
        iterator = ([row.sku, row.barcode_id, row.name, row.category, row.size or "", row.color or "", row.condition or "", int(row.in_stock_count or 0), int(row.selling_price or 0), row.inventory_status or "", local_date(row.created_at)] for row in query.yield_per(500))
    else:
        header = ["Staff ID", "Name", "Role", "Phone", "Active", "Created date (EAT)"]
        query = scoped(session.query(User), User.created_at, lower, upper).order_by(User.created_at, User.id)
        iterator = ([row.id, row.name, row.role, row.phone or "", "Yes" if row.active else "No", local_date(row.created_at)] for row in query.yield_per(500))
    return header, (row for row in iterator if matches(row, search))


def safe_text(value):
    # Excel / Sheets formula injection is distinct from CSV escaping. Preserve
    # numeric ledger amounts, but untrusted text must never execute as a formula.
    text = str(value if value is not None else "")
    if re.match(r"^[\s\ufeff]*[=+\-@]", text): text = "'" + text
    return text


def render_file(header, rows, file_format, title):
    count = 0
    if file_format == "csv":
        output = io.StringIO(newline=""); output.write("\ufeff")
        writer = csv.writer(output, lineterminator="\r\n"); writer.writerow(header)
        for row in rows:
            count += 1
            if count > MAX_ROWS: raise ValueError("Too many rows for one export. Choose a smaller date range.")
            writer.writerow([value if isinstance(value, (int, float)) else safe_text(value) for value in row])
            if output.tell() > MAX_BYTES: raise ValueError("The report is too large. Choose a smaller date range.")
        data = output.getvalue().encode("utf-8")
    else:
        workbook = Workbook(write_only=True); sheet = workbook.create_sheet(title[:31])
        sheet.freeze_panes = "A2"
        for index, label in enumerate(header, 1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(18, len(label) + 3))
        cells = []
        for label in header:
            cell = WriteOnlyCell(sheet, value=label); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="285447"); cells.append(cell)
        sheet.append(cells)
        try:
            for row in rows:
                count += 1
                if count > MAX_ROWS: raise ValueError("Too many rows for one export. Choose a smaller date range.")
                cells = []
                for value in row:
                    if isinstance(value, (int, float)): cell = WriteOnlyCell(sheet, value=value); cell.number_format = '#,##0'
                    else:
                        if len(str(value if value is not None else "")) > 32767:
                            raise ValueError("A record is too long for an Excel cell. Choose CSV to preserve the complete text.")
                        cell = WriteOnlyCell(sheet, value=re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value if value is not None else ""))); cell.data_type = "s"; cell.number_format = "@"
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{count + 1}"
            output = io.BytesIO(); workbook.save(output); data = output.getvalue()
        finally:
            if not sheet.closed: sheet.close()
            # Write-only worksheets stage XML on disk. Remove it even if a
            # row-limit / bad-cell error prevented workbook.save().
            writer = getattr(sheet, "_writer", None)
            if writer:
                try: writer.cleanup()
                except FileNotFoundError: pass
            workbook.close()
    if len(data) > MAX_BYTES: raise ValueError("The report is too large. Choose a smaller date range.")
    return data, count


def prepare(body, headers, staff, requires_staff_auth):
    try: kind, file_format, lower, upper, start, end, search = options(body)
    except ValueError as exc: return response({"ok": False, "error": str(exc)}, 400)
    now = datetime.utcnow()
    with get_db() as session:
        access = access_control.state(headers, session)
        area = KINDS[kind][1]
        if not access_control.allowed(area, access):
            status, result = access_control.locked_response(area); return response(result, status)
        if kind == "staff" and staff.get("role") not in {"admin", "manager"}:
            return response({"ok": False, "error": "An administrator or manager must export the staff directory."}, 403)
        header, rows = report_rows(session, kind, lower, upper, search)
        try: data, count = render_file(header, rows, file_format, KINDS[kind][0])
        except ValueError as exc: return response({"ok": False, "error": str(exc)}, 400)
        token = secrets.token_urlsafe(32)
        period = f"{start.isoformat() if start else 'beginning'}_to_{end.isoformat() if end else 'latest'}" if start or end else "all-dates"
        filename = f"adonai-{kind}-{period}-{now.replace(tzinfo=timezone.utc).astimezone(EAT):%Y%m%d-%H%M%S}.{file_format}"
        session.query(ExportDownload).filter(ExportDownload.expires_at <= now).delete(synchronize_session=False)
        # Bound abandoned capability storage; these are report copies, not books.
        old = session.query(ExportDownload).filter_by(staff_id=staff.get("id", "OPEN-TERMINAL")).order_by(ExportDownload.created_at.desc()).offset(19).all()
        for item in old: session.delete(item)
        item = ExportDownload(token_hash=hashlib.sha256(token.encode()).hexdigest(), filename=filename,
            mime_type=MIMES[file_format], kind=kind, record_count=count, byte_length=len(data), data=data,
            token_hashes_json=json.dumps([access_control._digest(value) for value in access_control._tokens(headers)]),
            staff_id=str(staff.get("id", "OPEN-TERMINAL")), requires_staff_auth=requires_staff_auth, expires_at=now + LIFETIME)
        session.add(item); session.flush()
        return response({"ok": True, "download": {"url": "/api/exports/download/" + token, "filename": filename,
            "type": item.mime_type, "bytes": len(data), "records": count, "kind": kind, "format": file_format,
            "from": start.isoformat() if start else None, "to": end.isoformat() if end else None,
            "timezone": "Africa/Kampala", "expires_at": item.expires_at.isoformat() + "Z"}}, 201)


def download(method, token, requires_staff_auth):
    if not re.fullmatch(r"[A-Za-z0-9_-]{43}", token): return response({"ok": False, "error": "Download not found."}, 404)
    with get_db() as session:
        item = session.get(ExportDownload, hashlib.sha256(token.encode()).hexdigest())
        if not item: return response({"ok": False, "error": "This download has expired. Export the report again."}, 410)
        if item.expires_at <= datetime.utcnow():
            session.delete(item); return response({"ok": False, "error": "This download has expired. Export the report again."}, 410)
        if requires_staff_auth and not item.requires_staff_auth:
            return response({"ok": False, "error": "Store protection changed. Export the report again."}, 403)
        if item.requires_staff_auth:
            user = session.get(User, item.staff_id)
            if user and not user.active:
                return response({"ok": False, "error": "This staff account is inactive. Export the report again with an active account."}, 403)
        access = access_control.state(session=session, token_hashes=json.loads(item.token_hashes_json))
        if not access_control.allowed(KINDS[item.kind][1], access):
            return response({"ok": False, "error": "This area has been locked. Unlock it and export the report again."}, 403)
        headers = {"Content-Disposition": f'attachment; filename="{item.filename}"; filename*=UTF-8\'\'{urllib.parse.quote(item.filename)}',
                   "Content-Length": str(item.byte_length), "Cache-Control": "private, no-store", "Referrer-Policy": "no-referrer"}
        return MediaResponse(200, item.mime_type, b"" if method == "HEAD" else bytes(item.data), headers)
